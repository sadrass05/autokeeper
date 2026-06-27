# -*- coding: utf-8 -*-
"""
生成 TRAE AI 创造力大赛报名帖 Word 文档
项目：自动记账助手（AutoBookkeeper）
赛道：生活娱乐
"""

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn

doc = Document()

# ============================================================
# 全局样式设置
# ============================================================
style = doc.styles['Normal']
font = style.font
font.name = '微软雅黑'
font.size = Pt(11)
style.element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')

# 页边距
for section in doc.sections:
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1.1)
    section.right_margin = Inches(1.1)


def add_heading_custom(text, level=1):
    """添加自定义标题"""
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.name = '微软雅黑'
        run.element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
        if level == 0:
            run.font.size = Pt(22)
            run.font.color.rgb = RGBColor(0x1a, 0x1a, 0x2e)
        elif level == 1:
            run.font.size = Pt(16)
            run.font.color.rgb = RGBColor(0x2b, 0x6c, 0xb0)
        elif level == 2:
            run.font.size = Pt(13)
            run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
    return h


def add_body(text, bold=False, italic=False):
    """添加正文段落"""
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = '微软雅黑'
    run.element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
    run.font.size = Pt(11)
    run.bold = bold
    run.italic = italic
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.5
    return p


def add_bullet(text, bold_prefix=None):
    """添加项目符号段落"""
    p = doc.add_paragraph(style='List Bullet')
    if bold_prefix:
        run_bold = p.add_run(bold_prefix)
        run_bold.bold = True
        run_bold.font.name = '微软雅黑'
        run_bold.element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
        run_bold.font.size = Pt(11)
    run = p.add_run(text)
    run.font.name = '微软雅黑'
    run.element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
    run.font.size = Pt(11)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.4
    return p


def add_label(label, content):
    """添加「标签：内容」格式的段落"""
    p = doc.add_paragraph()
    run_label = p.add_run(label)
    run_label.bold = True
    run_label.font.name = '微软雅黑'
    run_label.element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
    run_label.font.size = Pt(11)
    run_label.font.color.rgb = RGBColor(0x2b, 0x6c, 0xb0)
    run_content = p.add_run(content)
    run_content.font.name = '微软雅黑'
    run_content.element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
    run_content.font.size = Pt(11)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.5
    return p


def add_divider():
    """添加分隔线"""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run('─' * 40)
    run.font.color.rgb = RGBColor(0xcc, 0xcc, 0xcc)
    run.font.size = Pt(10)


# ============================================================
# 文档标题
# ============================================================
title = add_heading_custom('TRAE AI 创造力大赛 · 报名帖', level=0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run('赛道：生活娱乐  |  项目：自动记账助手 AutoBookkeeper')
run.font.name = '微软雅黑'
run.element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
run.font.size = Pt(12)
run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

add_divider()

# ============================================================
# 第一章：创意名称 + 创意介绍
# ============================================================
add_heading_custom('一、创意名称 + 创意介绍', level=1)

add_label('创意名称：', '自动记账助手（AutoBookkeeper）')

add_heading_custom('想解决什么问题', level=2)
add_body('"坚持记账不到两天"是绝大多数人的真实写照。每次消费完，掏出手机、打开记账 App、选分类、输金额、手动确认保存——这套流程繁琐到让人两三天就放弃。月底再翻支付宝、微信、云闪付的账单一笔笔补录，本质上是在做审计而不是记账。')
add_body('更麻烦的是，微信、支付宝、云闪付、美团、京东白条……支付渠道越来越多，数据散落在各个 App 里，没有一个地方能把它们合在一起看。而各品牌手机自带的钱包 App 打开就是铺天盖地的借贷入口和分期推广，根本不是为了帮你管钱而设计的。')

add_heading_custom('为什么会想到做这个', level=2)
add_body('后来我了解到 Android 有个 NotificationListenerService，可以合法读取手机的通知内容，不需要 root，不需要侵入其他 App。微信、支付宝、云闪付付完款后都会弹出一条通知，里面明明白白写着付给了谁、花了多少钱。既然如此，为什么不直接把这些通知抓下来，自动记进数据库？')
add_body('这个思路绕过了一个巨大的难题：不需要对接任何支付平台的 API，不需要操心微信的企业资质申请，不需要去谈支付宝的数据接口——只是读了一条本来就推送到手机上的通知而已。数据完全跑在本地，谁也看不到你的账单。')

add_heading_custom('大概是什么产品', level=2)
add_body('一款 Android App。监听微信、支付宝等 14 个平台的支付通知后自动识别并记录每一笔消费，数据完全本地存储，支持支出趋势可视化、分类饼图、CSV 备份，Pro 版还支持理财持仓管理和局域网数据同步。', bold=True)

add_divider()

# ============================================================
# 第二章：目标用户及痛点
# ============================================================
add_heading_custom('二、目标用户及痛点', level=1)

add_heading_custom('面向哪些用户', level=2)
add_bullet('想记账但总是坚持不下来的上班族和学生——手动记账的放弃率超过 90%', '记账困难户：')
add_bullet('同时使用微信、支付宝、云闪付、美团、京东等多个支付平台，账单散落各处无法统一查看的人', '多平台支付用户：')
add_bullet('反感借贷广告、对数据隐私敏感，不希望消费数据被拿去做用户画像的用户', '隐私敏感用户：')
add_bullet('希望统一管理基金、股票、定期存款持仓，并导出到 Excel 做统计分析的用户（Pro 版）', '理财管理用户：')

add_heading_custom('在什么场景下使用', level=2)
add_bullet('日常消费（早餐、打车、网购、缴费）后，App 在后台静默完成记账，用户完全无感', '日常消费：')
add_bullet('打开 App 查看本月支出总览、7 天趋势折线图、月度柱状图、分类环形饼图', '月底复盘：')
add_bullet('随时查看基金/股票/定期存款的持仓和收益排行，导出 CSV 到 Excel 分析', '理财管理：')
add_bullet('通过每周自动 CSV 备份和一键导入导出，确保换机或重装不丢数据', '换机备份：')

add_heading_custom('当前痛点', level=2)
add_body('如果没有这个产品，用户现在会遇到以下不便：')
add_bullet('每次消费后要打开 App、选分类、输金额、确认保存，超过 90% 的用户一周内放弃', '手动记账坚持不下来：')
add_bullet('微信、支付宝、云闪付、京东、美团账单互不相通，无法在一个地方统一查看，月底要挨个翻', '支付数据散落各处：')
add_bullet('手机自带钱包 App 查个消费记录满屏都是分期推广和借贷入口，每次都提心吊胆怕误点', '自带钱包变借贷入口：')
add_bullet('例如"用理财收益 cover 日常支出，不计入当月总支出"这种真实需求，通用 App 根本做不到', '通用 App 无法个性化：')
add_bullet('多数记账 App 会把消费数据上传云端用于用户画像分析，账单隐私无法保障', '数据安全焦虑：')

add_divider()

# ============================================================
# 第三章：技术创新亮点
# ============================================================
add_heading_custom('三、技术创新亮点', level=1)

add_heading_custom('创新点 1：通知监听「后门」方案——绕过支付平台 API 限制', level=2)
add_body('核心思路是利用 Android 系统内置的 NotificationListenerService，合法读取手机通知栏内容。微信、支付宝、云闪付等 App 在付款成功后会推送一条包含金额和商户信息的通知，这个通知本来就推送到用户手机上，App 只是把它抓下来自动解析入库。')
add_body('这个方案绕过了获取支付软件数据接口的难题：不需要企业资质申请微信支付 API，不需要谈支付宝数据接口，不需要 root 权限，也不需要侵入其他 App。同时通过前台服务 + 低优先级通知保活，断连时自动 requestRebind 重连，启动时还会扫描已有通知（最多 50 条）补录漏记记录。', bold=False)

add_heading_custom('创新点 2：14 平台智能解析引擎', level=2)
add_body('PaymentParser 模块支持 14 个支付平台的通知解析：')
add_bullet('微信、支付宝、拼多多、云闪付、美团、京东')
add_bullet('招商银行、工商银行、中国银行、建设银行、农业银行（5 大银行 App）')
add_bullet('抖音、快手、滴滴')
add_body('解析引擎包含三层智能提取：')
add_bullet('从"已支付¥12.50"到"一笔12.50元的支出"再到"消费88元"，5 级正则逐级降级匹配，金额范围校验 0.01~99999.99 元，确保不漏记也不误判', '5 级金额提取：')
add_bullet('从"付款给：XXX""向XXX支付""收款方：XXX""在XXX消费"等 9 种模式中提取商户名，长度和内容双重校验', '9 级商户名提取：')
add_bullet('自动识别零钱通、零钱、余额宝、余额、花呗、信用卡、借记卡、储蓄卡、银行卡等支付渠道。特别地，能识别拼多多"支付宝调用银行卡"这种复合渠道——检测到通知中同时包含"支付宝"和"银行卡"关键词时，标记为复合渠道而非简单银行卡', '智能渠道识别：')

add_heading_custom('创新点 3：国产 ROM 保活体系', level=2)
add_body('国产 ROM（MIUI、ColorOS、EMUI 等）会激进地杀死后台服务，导致通知监听中断。项目构建了一套完整的保活体系：')
add_bullet('通过反射调用隐藏 API SystemProperties.get，读取 ro.miui.ui.version.name、ro.hyperos.version、ro.coloros.version 等系统属性，精准识别 MIUI/HyperOS/ColorOS/EMUI/OriginOS/OneUI 六大 ROM', 'RomDetector 精准识别：')
add_bullet('针对不同 ROM 提供差异化引导——MIUI 需要最近任务锁定，ColorOS 需要关闭电池优化，EMUI 需要开启自启动，引导用户一步步设置', '差异化引导：')
add_bullet('NlsHealthState 心跳检测 + NlsWatchdogWorker 看门狗定时巡检 + NlsRestartWorker 自动重启，三重保障确保通知监听不被系统杀死', '三重保活机制：')

add_heading_custom('创新点 4：多层过滤防误判', level=2)
add_body('支付通知和消费统计、优惠推送、收款到账等通知混在一起，直接记录会产生大量垃圾数据。项目设计了四层过滤：')
add_bullet('"本周消费统计""月度账单""额度提醒""满减优惠"等 30+ 关键词，第一关拦截统计和营销通知', '第一层 · 黑名单过滤：')
add_bullet('"收款成功""退款到账""工资到账""红包到账"等 20+ 关键词，过滤收入类通知只记支出', '第二层 · 收入通知过滤：')
add_bullet('5 级正则全部未命中或金额超出 0.01~99999.99 范围则丢弃', '第三层 · 金额校验：')
add_bullet('基于通知的 sbn.key 生成唯一 notificationId，重复通知自动去重，已存在记录则更新字段而非重复插入', '第四层 · 防重复机制：')

add_heading_custom('创新点 5：Product Flavors 双版本架构', level=2)
add_body('利用 Android Gradle 的 Product Flavors 功能，在编译期将应用拆分为两个独立版本：')
add_bullet('纯记账功能，无理财代码，包名 com.example.autobookkeeper，3 个底部 Tab', 'Standard 版：')
add_bullet('在 Standard 全部功能基础上增加理财模块，包名 com.example.autobookkeeper.pro，4 个底部 Tab，支持基金/股票/定期持仓管理、收益排行、MySQL 局域网同步', 'Pro 版：')
add_body('两个版本包名不同，可同时安装在同一设备上。Standard 版通过 FlavorConfig 接口抽象实现零 Pro 残留代码，导入的理财数据自动转为普通支出记录。')

add_divider()

# ============================================================
# 第四章：AI Agent 协作开发故事
# ============================================================
add_heading_custom('四、AI Agent 协作开发故事', level=1)

add_heading_custom('一个不懂代码的人如何做出一个 App', level=2)
add_body('说实话，这个 App 几乎全是 AI Agent 写的。我不是程序员，完全不懂 Kotlin、Android 开发、Jetpack Compose 这些东西。我的工作流程大概是：把想法告诉 Trae → 它生成代码 → 我在 Android Studio 里跑起来 → 遇到报错就问 AI → 根据反馈不断改 → 一步步修到能跑为止。')

add_heading_custom('Trae AI Agent 的协作工作流', level=2)
add_body('整个开发过程中，Trae 承担了从架构设计到代码实现再到 Bug 修复的全链路工作：')
add_bullet('我描述"想要监听微信支付通知自动记账"，Trae 设计了 NotificationListenerService 方案并生成完整代码', '需求转代码：')
add_bullet('遇到 Hilt 与 Kotlin 2.1.x 不兼容、ViewModel 无法实例化闪退、Composable 函数不能直接存入列表等问题，Trae 逐个定位根因并给出修复方案', '踩坑与修复：')
add_bullet('从单版本到 Product Flavors 双版本架构、从基础记账到理财持仓管理，每次功能扩展都由 Trae 规划设计方案再逐步实现', '架构演进：')
add_bullet('项目积累了完整的单元测试，覆盖支付解析、ROM 检测、NLS 健康状态、备份恢复等核心逻辑', '测试覆盖：')

add_heading_custom('真实踩坑案例（摘自技术文档）', level=2)
add_body('以下是开发过程中 Trae 协助解决的真实技术问题：')
add_bullet('ViewModel 中使用 StateFlow 但 UI 层错误调用 LiveData 专用的 observeAsState，Trae 定位后改为 collectAsStateWithLifecycle', '问题一：')
add_bullet('Kotlin 2.1.x 的 Metadata 版本 2.1.0 超出旧版 Hilt 支持上限 2.0.0，Trae 建议升级 Hilt 到 2.56+', '问题二：')
add_bullet('MainViewModel 构造函数需要 Repository 参数但未接入 Hilt 注入体系，Trae 指导完成 ViewModel、Activity、Application 三处注解 + hiltViewModel 改造', '问题三：')
add_bullet('将 HomeScreen() 直接调用后存入列表返回 Unit 导致页面空白，Trae 指出需用 lambda 包裹实现延迟调用', '问题四：')

add_heading_custom('Agent 时代的意义', level=2)
add_body('Agent 时代的到来，确实给了普通人一次自定义工具的机会。放在以前，我这种没有任何编程背景的人想为自己做一个记账 App 简直是天方夜谭。虽然现在做出来的东西远不能跟大公司打磨多年的商业产品比——UI 不够精致、兼容性偶尔翻车、有些功能还只存在于规划里——但至少，我需要什么功能，我就加什么功能。不需要忍受广告，不需要付费订阅，数据不会被拿去分析用户画像。', bold=True)
add_body('这个项目本身就是 Agent 时代最好的注脚：一个零编程背景的普通用户，借助 Trae AI Agent，从零做出了一个包含通知监听、智能解析、ROM 保活、双版本架构、理财管理、数据同步的完整 Android 应用。')

add_divider()

# ============================================================
# 第五章：隐私与数据主权
# ============================================================
add_heading_custom('五、隐私与数据主权', level=1)

add_body('这个 App 的设计原则是数据主权归用户所有。在数据隐私日益重要的今天，这一点比功能本身更值得关注。')

add_heading_custom('数据 100% 本地存储', level=2)
add_body('所有消费记录存储在设备本地的 Room 数据库中，不上传任何远程服务器。通知读取权限仅用于识别支付消息，不会读取聊天内容。导入导出均为本地 CSV 文件，用户完全掌控数据的去向。')

add_heading_custom('Pro 版网络通信仅限局域网', level=2)
add_body('Pro 版支持将数据同步到自建 MySQL 数据库，但网络通信严格限制在局域网范围内。network_security_config.xml 默认禁止所有明文 HTTP 流量，仅放行私有 IP 段（10.x.x.x、192.168.x.x、172.16-31.x.x）。数据不会经过任何公网服务器中转。')

add_heading_custom('无广告、无追踪、无画像', level=2)
add_bullet('不含任何广告 SDK，不会推送借贷、分期、理财产品广告', '零广告：')
add_bullet('不含任何第三方统计 SDK（如友盟、Bugly），不收集设备信息、使用行为、崩溃日志', '零追踪：')
add_bullet('Release 包经 R8 全量混淆，逆向分析难度高，进一步保护用户数据安全', '代码混淆：')
add_bullet('消费数据不会被上传到云端做用户画像分析，你的账单只有你能看到', '零画像：')

add_heading_custom('与主流记账 App 的隐私对比', level=2)
add_body('主流记账 App 通常会：上传消费数据到云端、嵌入第三方统计 SDK、根据消费记录推送个性化广告、要求注册账号绑定手机号。而 AutoBookkeeper 以上四点全部不做——从设计层面就拒绝了数据变现的路径。', bold=True)

add_divider()

# ============================================================
# 第六章：价值与意义
# ============================================================
add_heading_custom('六、价值与意义', level=1)

add_heading_custom('效率提升：从「两天放弃」到「不知不觉就在记」', level=2)
add_body('传统记账 App 的核心矛盾是：记账本身是一件反人性的事——它要求用户在每次消费后主动做一系列操作。AutoBookkeeper 把"消费→记账"这个动作从"用户主动操作"变成"App 静默完成"，让记账这个好习惯真正能跑下去。')
add_body('用户不需要改变任何消费习惯，不需要记得打开 App，不需要选分类输金额——只需要正常用微信、支付宝付款，剩下的全部由 App 在后台自动完成。从"两天放弃"变成"不知不觉就在记"，这是效率层面的核心价值。')

add_heading_custom('社会价值：数据主权回归用户', level=2)
add_body('在数据隐私日益重要的今天，主流记账 App 普遍采用"免费使用 + 数据变现"的商业模式：用户的消费数据被上传到云端，用于用户画像分析、个性化广告推送、甚至金融产品推荐。用户在享受"免费"记账服务的同时，实际上在用自己最敏感的财务数据付费。')
add_body('AutoBookkeeper 从设计层面就拒绝了这条路径：数据 100% 本地存储，无广告无追踪无画像，Pro 版网络通信仅限局域网。它证明了一件事——一个不收集用户数据的记账工具，完全可以靠技术方案而非数据变现来满足用户需求。', bold=True)

add_heading_custom('社会价值：Agent 时代的个人工具样本', level=2)
add_body('这个项目更深层的社会价值在于：它是一个 Agent 时代的个人工具样本。一个零编程背景的普通用户，借助 Trae AI Agent，从零做出了一个包含通知监听、14 平台智能解析、国产 ROM 保活、双版本架构、理财管理、局域网数据同步的完整 Android 应用。')
add_body('放在以前，这种复杂度的 App需要专业开发团队数月的协作。而在 Agent 时代，一个有想法的普通人就能为自己量身定制工具——不需要忍受广告，不需要付费订阅，数据不会被拿去分析。这或许就是 Agent 时代最动人的图景：技术不再是少数人的特权，而是每个人表达需求的工具。', bold=True)

add_divider()

# ============================================================
# 技术栈附录
# ============================================================
add_heading_custom('附：技术栈一览', level=1)

tech_items = [
    ('UI 框架', 'Jetpack Compose + Material3'),
    ('架构模式', 'MVVM + Hilt 依赖注入'),
    ('本地数据库', 'Room + Flow 响应式查询'),
    ('网络请求', 'Retrofit2 + OkHttp3（Pro 版 MySQL 同步）'),
    ('图表可视化', 'MPAndroidChart'),
    ('后台任务', 'WorkManager（定时备份 + NLS 看门狗）'),
    ('构建变体', 'Product Flavors（standard / pro 编译期隔离）'),
    ('代码混淆', 'R8 / ProGuard 全量混淆'),
    ('开发方式', 'Trae AI Agent 协作开发（零编程背景）'),
]

for label, value in tech_items:
    add_bullet(value, f'{label}：')

# ============================================================
# 保存文档
# ============================================================
output_path = r'a:\awork\money\TRAE大赛报名帖_自动记账助手.docx'
doc.save(output_path)
print(f'文档已生成：{output_path}')
